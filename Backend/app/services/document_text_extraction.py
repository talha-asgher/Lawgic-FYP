"""Extract plain text from PDF, DOCX, TXT, and images; OCR when needed."""

from __future__ import annotations

import io
import logging
import os
from pathlib import PurePath

from app.services.document_ocr import (
    extract_pdf_text_pymupdf,
    ocr_image_bytes,
    ocr_pdf_bytes,
)

logger = logging.getLogger(__name__)

PDF_TYPES = frozenset({"application/pdf"})
DOCX_TYPES = frozenset(
    {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
)
TXT_TYPES = frozenset({"text/plain"})
IMAGE_TYPES = frozenset(
    {
        "image/png",
        "image/jpeg",
        "image/jpg",
        "image/webp",
    }
)

MIN_PDF_TEXT_CHARS = max(0, int(os.environ.get("DOC_ANALYSIS_MIN_PDF_TEXT_CHARS", "50")))


def _suffix(name: str) -> str:
    return PurePath(name or "").suffix.lower()


def detect_kind(file_name: str, content_type: str | None) -> str | None:
    """Return 'pdf', 'docx', 'txt', 'image', or None if unsupported."""
    ct = (content_type or "").split(";")[0].strip().lower()
    suf = _suffix(file_name)

    if ct in PDF_TYPES or suf == ".pdf":
        return "pdf"
    if ct in DOCX_TYPES or suf == ".docx":
        return "docx"
    if ct in TXT_TYPES or suf == ".txt":
        return "txt"
    if ct in IMAGE_TYPES or suf in (".png", ".jpg", ".jpeg", ".webp"):
        return "image"
    return None


def extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        t = page.extract_text()
        if t and t.strip():
            parts.append(t.strip())
    return "\n\n".join(parts)


def extract_pdf_with_fallbacks(data: bytes) -> str:
    """Embedded text first; if too short, PyMuPDF text; then Tesseract page OCR."""
    text = extract_pdf(data).strip()
    if len(text) >= MIN_PDF_TEXT_CHARS:
        return text
    alt = extract_pdf_text_pymupdf(data).strip()
    if len(alt) > len(text):
        text = alt
    if len(text) >= MIN_PDF_TEXT_CHARS:
        return text
    ocr = ocr_pdf_bytes(data).strip()
    if ocr:
        return ocr
    return text or ocr


def extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for p in doc.paragraphs:
        if p.text and p.text.strip():
            lines.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_document_text(data: bytes, file_name: str, content_type: str | None) -> str:
    kind = detect_kind(file_name, content_type)
    if kind is None:
        raise ValueError(
            "Unsupported file type. Upload a PDF, DOCX, TXT, or image (PNG, JPEG, WebP)."
        )
    if kind == "pdf":
        text = extract_pdf_with_fallbacks(data)
    elif kind == "docx":
        text = extract_docx(data)
    elif kind == "image":
        text = ocr_image_bytes(data)
    else:
        text = extract_txt(data)
    text = " ".join(text.split())
    if len(text) > 200_000:
        logger.info("Truncating extracted text from %s to 200000 chars", len(text))
        text = text[:200_000]
    return text
